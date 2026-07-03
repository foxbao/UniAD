from __future__ import annotations

import csv
import json
import pickle
import re
import shutil
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from PIL import Image, ImageDraw, ImageFont


BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parents[2]
FIG_DIR = BASE_DIR / "figures"
SRC_DIR = BASE_DIR / "sources"
OUT_MD = BASE_DIR / "turnaware_motion_patent_disclosure_v0.5_loss_anchor_combined.md"
OUT_DOCX = BASE_DIR / "turnaware_motion_patent_disclosure_v0.5_loss_anchor_combined.docx"
OUT_REPRO_MD = BASE_DIR / "turnaware_motion_patent_disclosure_v0.5_internal_repro.md"
REPRO_JSON = SRC_DIR / "v0.5_reproducibility_record.json"
FIG3_CASES = [
    (
        "fig3a_turn_focus_truck_1522.png",
        SRC_DIR
        / "fig3_focus_patent_ttl_zh"
        / "000_001522_1535527a-6460-4634-9a9a-5f909833df1f.png",
    ),
    (
        "fig3b_turn_focus_crane_204.png",
        SRC_DIR
        / "fig3_focus_patent_ttl_zh"
        / "001_000204_9a8da77a-295b-41c9-8e1e-3b218e2143c8.png",
    ),
    (
        "fig3c_turn_focus_crane_2409.png",
        SRC_DIR
        / "fig3_focus_patent_ttl_zh"
        / "002_002409_15f680cc-82b9-4950-b77c-c14ea1686465.png",
    ),
]

TITLE = "一种面向港口自动驾驶车辆的转弯感知多模态运动预测方法、装置、设备及存储介质"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def draw_centered(draw: ImageDraw.ImageDraw, xy, text: str, fnt, fill="#111827") -> None:
    x0, y0, x1, y1 = xy
    if hasattr(draw, "multiline_textbbox"):
        bbox = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=4, align="center")
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    else:
        lines = text.splitlines() or [text]
        sizes = [draw.textsize(line, font=fnt) for line in lines]
        tw = max((size[0] for size in sizes), default=0)
        th = sum(size[1] for size in sizes) + max(0, len(lines) - 1) * 4
    draw.multiline_text(
        (x0 + (x1 - x0 - tw) / 2, y0 + (y1 - y0 - th) / 2),
        text,
        font=fnt,
        fill=fill,
        spacing=4,
        align="center",
    )


def rounded_box(draw, xy, fill, outline="#64748B", radius=16, width=2) -> None:
    if hasattr(draw, "rounded_rectangle"):
        draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)
    else:
        draw.rectangle(xy, fill=fill, outline=outline)
        if width > 1:
            for offset in range(1, width):
                draw.rectangle(
                    (xy[0] + offset, xy[1] + offset, xy[2] - offset, xy[3] - offset),
                    outline=outline,
                )


def arrow(draw, start, end, fill="#334155", width=4) -> None:
    draw.line([start, end], fill=fill, width=width)
    x0, y0 = start
    x1, y1 = end
    if abs(x1 - x0) >= abs(y1 - y0):
        sign = 1 if x1 > x0 else -1
        pts = [(x1, y1), (x1 - sign * 16, y1 - 8), (x1 - sign * 16, y1 + 8)]
    else:
        sign = 1 if y1 > y0 else -1
        pts = [(x1, y1), (x1 - 8, y1 - sign * 16), (x1 + 8, y1 - sign * 16)]
    draw.polygon(pts, fill=fill)


def poly_arrow(draw, points, fill="#334155", width=4) -> None:
    draw.line(points, fill=fill, width=width, joint="curve")
    x0, y0 = points[-2]
    x1, y1 = points[-1]
    if abs(x1 - x0) >= abs(y1 - y0):
        sign = 1 if x1 > x0 else -1
        pts = [(x1, y1), (x1 - sign * 16, y1 - 8), (x1 - sign * 16, y1 + 8)]
    else:
        sign = 1 if y1 > y0 else -1
        pts = [(x1, y1), (x1 - 8, y1 - sign * 16), (x1 + 8, y1 - sign * 16)]
    draw.polygon(pts, fill=fill)


def heading_change_deg(traj: np.ndarray) -> float:
    if traj.shape[0] < 3:
        return 0.0
    v0 = traj[1] - traj[0]
    v1 = traj[-1] - traj[-2]
    if np.linalg.norm(v0) < 1e-3 or np.linalg.norm(v1) < 1e-3:
        return 0.0
    angle = np.degrees(np.arctan2(v1[1], v1[0]) - np.arctan2(v0[1], v0[0]))
    return float((angle + 180.0) % 360.0 - 180.0)


def draw_text_centered_at(draw: ImageDraw.ImageDraw, center, text: str, fnt, fill="#111827") -> None:
    if hasattr(draw, "textbbox"):
        bbox = draw.textbbox((0, 0), text, font=fnt)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    else:
        tw, th = draw.textsize(text, font=fnt)
    draw.text((center[0] - tw / 2, center[1] - th / 2), text, font=fnt, fill=fill)


def draw_rotated_label(img: Image.Image, center, text: str, fnt, fill="#111827") -> None:
    layer = Image.new("RGBA", (260, 44), (255, 255, 255, 0))
    layer_draw = ImageDraw.Draw(layer)
    layer_draw.text((0, 0), text, font=fnt, fill=fill)
    bbox = layer.getbbox()
    if bbox:
        layer = layer.crop(bbox)
    rotated = layer.rotate(90, expand=True)
    xy = (int(center[0] - rotated.size[0] / 2), int(center[1] - rotated.size[1] / 2))
    img.paste(rotated, xy, rotated)


def draw_source_anchor_legend(draw: ImageDraw.ImageDraw, xy, anchors: np.ndarray, fnt) -> None:
    colors = ["#1F77B4", "#FF7F0E", "#2CA02C", "#D62728", "#9467BD", "#8C564B"]
    x0, y0 = xy
    box_w, box_h = 158, 136
    draw.rectangle((x0, y0, x0 + box_w, y0 + box_h), fill="#FFFFFF", outline="#D1D5DB", width=1)
    for idx, color in enumerate(colors[: len(anchors)]):
        y = y0 + 12 + idx * 20
        deg = heading_change_deg(anchors[idx])
        label = f"模式{idx}：{deg:+.0f}°"
        draw.line([(x0 + 10, y + 8), (x0 + 34, y + 8)], fill=color, width=3)
        draw.ellipse((x0 + 20, y + 4, x0 + 28, y + 12), fill=color)
        draw.text((x0 + 42, y - 2), label, font=fnt, fill="#111827")


def draw_source_anchor_panel(
    img: Image.Image,
    slot,
    title: str,
    anchors: np.ndarray,
    xlim,
    ylim,
    xticks,
    yticks,
    title_font,
    label_font,
    tick_font,
    legend_font,
) -> None:
    draw = ImageDraw.Draw(img)
    x0, y0, x1, y1 = slot
    draw_text_centered_at(draw, ((x0 + x1) / 2, y0 + 25), title, title_font)

    max_rect = (x0 + 64, y0 + 68, x1 - 28, y1 - 72)
    data_w = float(xlim[1] - xlim[0])
    data_h = float(ylim[1] - ylim[0])
    scale = min((max_rect[2] - max_rect[0]) / data_w, (max_rect[3] - max_rect[1]) / data_h)
    plot_w = data_w * scale
    plot_h = data_h * scale
    px0 = int(max_rect[0] + ((max_rect[2] - max_rect[0]) - plot_w) / 2)
    py0 = int(max_rect[1] + ((max_rect[3] - max_rect[1]) - plot_h) / 2)
    px1 = int(px0 + plot_w)
    py1 = int(py0 + plot_h)

    def map_pt(pt):
        return (
            int(round(px0 + (float(pt[0]) - xlim[0]) * scale)),
            int(round(py1 - (float(pt[1]) - ylim[0]) * scale)),
        )

    draw.rectangle((px0, py0, px1, py1), fill="#FFFFFF", outline="#111827", width=2)
    for tx in xticks:
        if xlim[0] <= tx <= xlim[1]:
            px, _ = map_pt((tx, ylim[0]))
            draw.line([(px, py0), (px, py1)], fill="#E5E7EB", width=1)
            draw.line([(px, py1), (px, py1 + 5)], fill="#111827", width=1)
            draw_text_centered_at(draw, (px, py1 + 18), str(tx), tick_font)
    for ty in yticks:
        if ylim[0] <= ty <= ylim[1]:
            _, py = map_pt((xlim[0], ty))
            draw.line([(px0, py), (px1, py)], fill="#E5E7EB", width=1)
            draw.line([(px0 - 5, py), (px0, py)], fill="#111827", width=1)
            label = str(ty)
            if hasattr(draw, "textbbox"):
                bbox = draw.textbbox((0, 0), label, font=tick_font)
                tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
            else:
                tw, th = draw.textsize(label, font=tick_font)
            draw.text((px0 - tw - 10, py - th / 2), label, font=tick_font, fill="#111827")
    if xlim[0] <= 0 <= xlim[1]:
        px, _ = map_pt((0, ylim[0]))
        draw.line([(px, py0), (px, py1)], fill="#6B7280", width=1)
    if ylim[0] <= 0 <= ylim[1]:
        _, py = map_pt((xlim[0], 0))
        draw.line([(px0, py), (px1, py)], fill="#6B7280", width=1)

    colors = ["#1F77B4", "#FF7F0E", "#2CA02C", "#D62728", "#9467BD", "#8C564B"]
    for traj, color in zip(anchors, colors):
        pts = np.vstack([np.zeros((1, 2), dtype=np.float32), traj])
        mapped = [map_pt(pt) for pt in pts]
        draw.line(mapped, fill=color, width=4, joint="curve")
        for mx, my in mapped:
            draw.ellipse((mx - 4, my - 4, mx + 4, my + 4), fill=color)
        ex, ey = mapped[-1]
        draw.ellipse((ex - 8, ey - 8, ex + 8, ey + 8), fill=color)

    draw_text_centered_at(draw, ((px0 + px1) / 2, py1 + 46), "x（横向，米）", label_font)
    draw_rotated_label(img, (px0 - 48, (py0 + py1) / 2), "y（前向，米）", label_font)
    draw_source_anchor_legend(draw, (min(px1 - 166, x1 - 188), py0 + 10), anchors, legend_font)


def generate_fig2_anchor_compare() -> None:
    with (SRC_DIR / "motion_anchor_infos_kl.pkl").open("rb") as f:
        old = pickle.load(f)
    with (SRC_DIR / "motion_anchor_infos_kl_turnaware_3grp.pkl").open("rb") as f:
        new = pickle.load(f)
    old_anchors = old["anchors_all"]
    new_anchors = new["anchors_all"]

    img = Image.new("RGB", (1500, 1800), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    main_font = font(30, bold=True)
    title_font = font(25, bold=True)
    label_font = font(20)
    tick_font = font(17)
    legend_font = font(16)
    draw_text_centered_at(
        draw,
        (750, 36),
        "运动锚点：普通锚点与转弯感知锚点对比（+Y为目标前向，原点为目标当前位置）",
        main_font,
    )
    panels = [
        ((60, 78, 710, 820), "行人：普通锚点（偏直行）", old_anchors[0], (-3, 3), (-5, 8), [-2, -1, 0, 1, 2], [-4, -2, 0, 2, 4, 6]),
        ((790, 78, 1440, 820), "行人：转弯感知锚点", new_anchors[0], (-4, 4.5), (-2.5, 7), [-3, -2, -1, 0, 1, 2, 3, 4], [-2, 0, 2, 4, 6]),
        ((60, 900, 710, 1725), "车辆/作业设备：普通锚点（偏直行）", old_anchors[1], (-5, 5), (-2, 38), [-2.5, 0, 2.5], [0, 5, 10, 15, 20, 25, 30, 35]),
        ((790, 900, 1440, 1725), "车辆/作业设备：转弯感知锚点", new_anchors[1], (-15, 13), (-2, 38), [-10, -5, 0, 5, 10], [0, 5, 10, 15, 20, 25, 30, 35]),
    ]
    for panel in panels:
        draw_source_anchor_panel(
            img,
            panel[0],
            panel[1],
            np.asarray(panel[2], dtype=np.float32),
            panel[3],
            panel[4],
            panel[5],
            panel[6],
            title_font,
            label_font,
            tick_font,
            legend_font,
        )
    img.save(FIG_DIR / "fig2_anchor_compare.png")


def generate_figures() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    title_font = font(28, bold=True)
    body_font = font(20)
    small_font = font(16)

    # Fig. 1: restrained patent-style flow chart in Chinese.
    fig1_body_font = font(24)
    fig1_small_font = font(18)
    img = Image.new("RGB", (1500, 980), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    flow_boxes = [
        ((475, 60, 1025, 145), "港口训练数据\n目标类别、航向角、未来轨迹"),
        ((475, 195, 1025, 280), "按目标类别分组\n行人、车辆/作业设备、锥桶"),
        ((475, 330, 1025, 415), "目标局部坐标归一\n统一目标朝向"),
        ((475, 465, 1025, 550), "计算运动形态特征\n路径、净位移、航向变化、横移比"),
        ((110, 645, 430, 745), "运动分桶\n静止低速、直行、转弯"),
        ((590, 645, 910, 745), "分层保留槽位\n短直行、长直行、左右转弯"),
        ((1070, 645, 1390, 745), "转弯样本加权\n提高长尾样本影响"),
        ((475, 850, 1025, 935), "生成多模态运动锚点\n输入运动预测网络"),
    ]
    for xy, text in flow_boxes:
        draw.rectangle(xy, fill="#FFFFFF", outline="#111827", width=3)
        draw_centered(draw, xy, text, fig1_body_font, fill="#111827")
    for points in [
        [(750, 145), (750, 195)],
        [(750, 280), (750, 330)],
        [(750, 415), (750, 465)],
        [(750, 550), (750, 600), (270, 600), (270, 645)],
        [(750, 550), (750, 645)],
        [(750, 550), (750, 600), (1230, 600), (1230, 645)],
        [(270, 745), (270, 805), (750, 805), (750, 850)],
        [(750, 745), (750, 850)],
        [(1230, 745), (1230, 805), (750, 805), (750, 850)],
    ]:
        poly_arrow(draw, points, fill="#111827", width=3)
    draw.text((90, 948), "说明：上方为主流程，中间三项为并列处理步骤，下方为锚点输出。", font=fig1_small_font, fill="#374151")
    img.save(FIG_DIR / "fig1_turnaware_anchor_generation_flow.png")

    # Fig. 2: directly render Chinese labels from original anchor pkl files.
    generate_fig2_anchor_compare()

    # Fig. 3A-3C: copy regenerated Chinese focused visual comparisons if present.
    for filename, source_path in FIG3_CASES:
        if source_path.exists():
            shutil.copyfile(source_path, FIG_DIR / filename)

    # Fig. 4: turn-aware loss and mode assignment flow.
    fig4_body_font = font(24)
    img = Image.new("RGB", (1500, 560), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    boxes = [
        ((80, 80, 360, 190), "真实未来轨迹\n有效掩码"),
        ((450, 80, 730, 190), "计算平均误差\n与终点误差"),
        ((820, 80, 1120, 190), "联合误差\n选择监督模式"),
        ((1200, 80, 1440, 190), "监督模式\n分类+回归"),
        ((80, 330, 360, 440), "路径/净位移\n航向变化/横移比"),
        ((450, 330, 730, 440), "运动分桶\nstatic/straight\nmild/sharp"),
        ((820, 330, 1120, 440), "样本权重\n按运动桶提高"),
        ((1200, 330, 1440, 440), "加权损失\n按批次均值归一"),
    ]
    for xy, text in boxes:
        rounded_box(draw, xy, "#FFFFFF", outline="#111827", radius=12, width=3)
        draw_centered(draw, xy, text, fig4_body_font)
    for start, end in [((360, 135), (450, 135)), ((730, 135), (820, 135)), ((1120, 135), (1200, 135)),
                       ((360, 385), (450, 385)), ((730, 385), (820, 385)), ((1120, 385), (1200, 385)),
                       ((1320, 330), (1320, 190)), ((960, 330), (960, 190))]:
        arrow(draw, start, end, fill="#111827", width=3)
    img.save(FIG_DIR / "fig4_turnaware_loss_mode_assignment_flow.png")

    # Fig. 5: highest-confidence / best-candidate diagnosis.
    fig5_body_font = font(24)
    fig5_small_font = font(20)
    img = Image.new("RGB", (1500, 820), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    # coordinate sketch
    draw.line([(170, 620), (900, 620)], fill="#111827", width=2)
    draw.line([(170, 620), (170, 130)], fill="#111827", width=2)
    origin = (190, 590)
    gt = [(190, 590), (270, 530), (360, 460), (480, 390), (620, 330), (800, 300)]
    top1 = [(190, 590), (300, 560), (430, 540), (570, 530), (720, 535), (880, 555)]
    oracle = [(190, 590), (260, 525), (350, 450), (470, 375), (610, 320), (780, 290)]
    other = [(190, 590), (240, 610), (300, 640), (370, 665), (450, 685), (540, 700)]
    for pts, color, label, text_offset in [
        (other, "#94A3B8", "其他候选", (-40, 30)),
        (top1, "#EF4444", "最高置信度候选：偏直", (-85, -60)),
        (oracle, "#2563EB", "最优候选：接近真实轨迹", (-60, -28)),
        (gt, "#16A34A", "真实转弯轨迹", (-40, 45)),
    ]:
        draw.line(pts, fill=color, width=6)
        for p in pts[1:]:
            draw.ellipse((p[0]-5, p[1]-5, p[0]+5, p[1]+5), fill=color)
        draw.text((pts[-1][0] + text_offset[0], pts[-1][1] + text_offset[1]), label, font=fig5_small_font, fill=color)
    draw.ellipse((origin[0]-12, origin[1]-12, origin[0]+12, origin[1]+12), fill="#111827")
    draw.text((130, 645), "目标当前位置", font=fig5_small_font, fill="#111827")
    # diagnosis table
    x0, y0 = 1040, 105
    rounded_box(draw, (x0, y0, 1430, 270), "#FFFFFF", outline="#111827", width=3)
    draw_centered(draw, (x0, y0, 1430, 270), "最优候选好\n但最高置信度候选差\n=> 模式评分未选中", fig5_body_font, fill="#111827")
    rounded_box(draw, (x0, 320, 1430, 485), "#FFFFFF", outline="#111827", width=3)
    draw_centered(draw, (x0, 320, 1430, 485), "最优候选和最高置信度\n候选都差\n=> 候选覆盖不足", fig5_body_font, fill="#111827")
    rounded_box(draw, (x0, 535, 1430, 700), "#FFFFFF", outline="#111827", width=3)
    draw_centered(draw, (x0, 535, 1430, 700), "未匹配目标\n=> 预测模块未获得对象\n=> 检测/跟踪问题", fig5_body_font, fill="#111827")
    img.save(FIG_DIR / "fig5_candidate_confidence_diagnosis.png")

    # Fig. 6: generic prediction network architecture and insertion points.
    fig6_font = font(25)
    img = Image.new("RGB", (1800, 820), "#FFFFFF")
    draw = ImageDraw.Draw(img)

    def flow_box(xy, text: str) -> None:
        draw.rectangle(xy, fill="#FFFFFF", outline="#111827", width=3)
        draw_centered(draw, xy, text, fig6_font, fill="#111827")

    boxes = {
        "input": ((80, 90, 330, 200), "传感器特征\n历史目标状态"),
        "encode": ((390, 90, 640, 200), "目标/场景\n特征编码"),
        "interact": ((700, 90, 950, 200), "场景交互编码\n目标与环境关系"),
        "decode": ((1010, 90, 1280, 200), "多模态轨迹解码器\n融合锚点先验"),
        "candidate": ((1340, 90, 1580, 200), "候选轨迹与置信度\nK条轨迹、K个分数"),
        "infer": ((1630, 90, 1770, 200), "推理输出\n前K条轨迹"),
        "gt": ((1010, 320, 1280, 440), "真实未来轨迹\n运动桶标签"),
        "loss": ((1340, 320, 1580, 440), "训练损失\n联合模式分配\n运动桶加权"),
        "target": ((80, 620, 330, 730), "目标类别、中心\n航向角"),
        "anchor": ((390, 620, 640, 730), "转弯感知锚点库\n类别组与运动族槽位"),
        "transform": ((700, 620, 950, 730), "锚点选择与坐标变换\n局部锚点到场景坐标"),
    }
    for xy, label in boxes.values():
        flow_box(xy, label)

    for start, end in [
        ((330, 145), (390, 145)),
        ((640, 145), (700, 145)),
        ((950, 145), (1010, 145)),
        ((1280, 145), (1340, 145)),
        ((1580, 145), (1630, 145)),
        ((1460, 200), (1460, 320)),
        ((330, 675), (390, 675)),
        ((640, 675), (700, 675)),
        ((1280, 380), (1340, 380)),
    ]:
        arrow(draw, start, end, fill="#111827", width=3)

    poly_arrow(draw, [(950, 675), (980, 675), (980, 250), (1145, 250), (1145, 200)], fill="#111827", width=3)
    draw.text((1010, 225), "锚点先验", font=font(21), fill="#111827")
    draw.text((1480, 250), "训练", font=font(21), fill="#111827")
    img.save(FIG_DIR / "fig6_prediction_network_architecture.png")


def load_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="") as f:
        return list(csv.DictReader(f))


def format_float(value: str, digits: int = 4) -> str:
    try:
        return f"{float(value):.{digits}f}"
    except Exception:
        return value


def table_md(headers: list[str], rows: list[list[str]]) -> str:
    out = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    out += ["| " + " | ".join(map(str, row)) + " |" for row in rows]
    return "\n".join(out)


def build_markdown() -> str:
    overall = load_csv_rows(SRC_DIR / "v0.5_epoch6_overall_ablation_summary.csv")
    turn_rows = load_csv_rows(SRC_DIR / "v0.5_turn_bucket_epoch6_summary.csv")
    def turn(model, bucket, key):
        for r in turn_rows:
            if r["model"] == model and r["bucket"] == bucket:
                return format_float(r[key])
        return ""

    config_display = {
        "base_e2e_lidar": "实验组A：普通锚点+原始训练目标",
        "base_e2e_lidar_turnaware": "实验组B：转弯感知锚点+原始训练目标",
        "base_e2e_lidar_turnloss": "实验组C：普通锚点+转弯分桶训练目标",
        "base_e2e_lidar_turnaware_turnloss": "实验组D：转弯感知锚点+转弯分桶训练目标",
        "base_e2e_lidar_turnaware_modescore": "实验组E：组合方案+转弯模式评分增强",
    }
    anchor_display = {
        "old 2grp": "普通类别锚点",
        "turn-aware 3grp": "转弯感知分组锚点",
    }
    loss_display = {
        "original": "原始训练目标",
        "ADE+FDE + turn weighting": "平均位移误差/终点位移误差联合分配+转弯加权",
    }
    group_label = {
        "base_e2e_lidar": "A",
        "base_e2e_lidar_turnaware": "B",
        "base_e2e_lidar_turnloss": "C",
        "base_e2e_lidar_turnaware_turnloss": "D",
        "base_e2e_lidar_turnaware_modescore": "E",
    }

    def overall_row(model: str) -> dict[str, str]:
        for row in overall:
            if row["config"] == model:
                return row
        return {}

    def metric(row: dict[str, str], key: str, digits: int = 4) -> str:
        return format_float(row.get(key, ""), digits=digits)

    def lower_better_change(base_value: str, value: str) -> str:
        try:
            base_float = float(base_value)
            value_float = float(value)
        except Exception:
            return ""
        if abs(base_float) < 1e-12:
            return ""
        pct = (base_float - value_float) / base_float * 100.0
        if abs(pct) < 0.05:
            return "基本持平"
        if pct > 0:
            return f"改善{pct:.1f}%"
        return f"退化{abs(pct):.1f}%"

    experiment_setup_table = table_md(
        ["实验组", "锚点方案", "训练目标", "用途说明"],
        [
            ["A", "普通类别锚点", "原始训练目标", "基线方案"],
            ["B", "转弯感知分组锚点", "原始训练目标", "验证锚点方案的独立贡献"],
            ["C", "普通类别锚点", "联合模式分配+转弯加权", "验证训练目标的独立贡献"],
            ["D", "转弯感知分组锚点", "联合模式分配+转弯加权", "验证锚点与训练目标的组合效果"],
            ["E", "转弯感知分组锚点", "联合模式分配+转弯加权+模式评分增强", "诊断体系导出的扩展实验"],
        ],
    )

    overall_summary_table = table_md(
        ["实验组", "minADE", "minFDE", "MR", "AMOTA", "NDS", "主要观察"],
        [
            [
                group_label.get(r["config"], r["config"]),
                metric(r, "minADE"),
                metric(r, "minFDE"),
                metric(r, "MR"),
                metric(r, "AMOTA"),
                metric(r, "NDS"),
                {
                    "base_e2e_lidar": "基线",
                    "base_e2e_lidar_turnaware": "锚点单独替换后终点误差略降",
                    "base_e2e_lidar_turnloss": "运动预测单项指标最优",
                    "base_e2e_lidar_turnaware_turnloss": "综合感知跟踪指标较好",
                }.get(r["config"], ""),
            ]
            for r in overall
        ],
    )

    baseline_model = "base_e2e_lidar"
    baseline_mild_fde = turn(baseline_model, "mild_turn", "minFDE")
    baseline_sharp_fde = turn(baseline_model, "sharp_turn", "minFDE")
    baseline_sharp_mr = turn(baseline_model, "sharp_turn", "MR")
    turn_focus_table = table_md(
        ["实验组", "轻微转弯minFDE", "较A变化", "大转弯minFDE", "较A变化", "大转弯MR", "较A变化"],
        [
            [
                group_label.get(m, m),
                turn(m, "mild_turn", "minFDE"),
                "基线" if m == baseline_model else lower_better_change(baseline_mild_fde, turn(m, "mild_turn", "minFDE")),
                turn(m, "sharp_turn", "minFDE"),
                "基线" if m == baseline_model else lower_better_change(baseline_sharp_fde, turn(m, "sharp_turn", "minFDE")),
                turn(m, "sharp_turn", "MR"),
                "基线" if m == baseline_model else lower_better_change(baseline_sharp_mr, turn(m, "sharp_turn", "MR")),
            ]
            for m, _ in turn_models_for_doc()
        ],
    )
    text = f"""# 坤浪科技专利技术交底书

## 基本信息

{table_md(
    ["项目", "内容"],
    [
        ["专利名称", TITLE],
        ["专利类型", "发明专利"],
        ["技术领域", "港口自动驾驶；场区智能感知；多目标运动预测；端到端感知预测模型"],
        ["发明人", "鲍佳立、贺志国、聂新勇"],
        ["联系人", "鲍佳立"],
        ["联系电话", "18117553298"],
        ["电子邮箱", "jiali.bao@kunlang.cn"],
    ],
)}

> 本文件为内部技术交底草案，用于与专利代理人沟通发明构思、实施方式、实验数据和可保护点。正式申请前建议由专利代理人进行新颖性/创造性检索、权利要求布局和公开风险审查。

## 创新点摘要

针对港口自动驾驶场景中集卡、拖挂车、智能导引运输车、吊车、叉车等目标存在大量静止、直行和低速样本，而真实转弯、掉头、横移、装卸作业等运动模式呈长尾分布，导致普通运动锚点偏向直线、运动预测在转弯场景表现较弱的问题，提出一种“目标局部坐标系下的统一转弯度量 + 港口类别分组 + 分层转弯锚点 + 转弯分桶训练损失”的多模态运动预测方案。

该方案先在目标局部坐标系下计算未来轨迹的路径长度、净位移、航向变化和横向偏移比例，将轨迹划分为静止/低速、短直行、长直行、左/右转弯和大机动转弯等运动族；再为不同类别组预留运动族锚点槽位，必要时结合转弯样本加权聚类生成各槽位的代表轨迹；训练时采用平均位移误差与终点位移误差联合的监督模式分配，并按转弯桶对回归损失和/或模式分类损失进行加权。该方案不依赖增加预测模式数量，可在推理复杂度基本不变的情况下改善港口车辆转弯候选覆盖和转弯轨迹训练权重。

## 一、背景技术

港口自动驾驶车辆通常在堆场、闸口、岸桥、轨道吊、泊位等区域运行。与城市道路车辆相比，港口目标类别更复杂，包括集卡、拖挂车、IGV 空/满载、轮胎吊、正面吊、叉车、行人、锥桶等；运动模式也更加复杂，既存在较长距离直行运输，也存在低速倒车、横移、原地转向、装卸作业、人工遥控导致的非规则转弯等行为。

现有多模态运动预测方法通常从训练集未来轨迹中直接进行普通聚类，得到固定数量的轨迹锚点，并用平均位移误差或类似指标选择监督模式。该类方法在普通道路数据集中可以覆盖部分直行、变道和转弯模式，但在港口场景中存在以下问题：

1. 港口训练集中静止、慢速和直行样本占比高，普通聚类容易被高频模式主导，车辆锚点缺少真正大角度转弯或作业式转向模板。
2. 真实转弯样本属于长尾样本，即便其对安全和规划重要，也容易在锚点生成和损失训练中被低估。
3. 仅增加预测模式数量会增加模型参数和推理成本，并且不能保证新增模式覆盖转弯。
4. 锥桶等近静止目标若与车辆或行人混合生成锚点，可能污染车辆/行人的运动锚点。
5. 多模态输出中经常出现“最优候选误差较小但最高置信度候选误差较大”的现象，说明候选轨迹已经存在，但模型置信度没有把正确预测模式排到第一。

与通用轨迹预测方法相比，本发明不以单纯增加候选轨迹数量、单纯改用带权聚类或单纯调整训练误差为核心，而是把港口类别分组、目标局部坐标系下的统一转弯度量、运动族槽位保留和转弯分桶训练目标组合使用。该组合使锚点结构与港口作业运动语义对应，并使训练过程显式关注转弯长尾样本，从而区别于仅依据全局轨迹距离聚类或仅依据平均误差选择监督模式的现有方案。

## 二、发明内容

### 1. 要解决的技术问题

本发明拟解决港口自动驾驶运动预测模型在转弯场景中预测能力不足的问题，尤其是：

1. 运动锚点偏直线，缺少转弯候选；
2. 训练损失被直行/静止样本主导，转弯样本权重不足；
3. 监督模式仅按平均位移误差选择时，容易忽略终点误差；
4. 候选轨迹存在但模式分类分数选错，导致实际推理时最高置信度轨迹偏差较大；
5. 港口异构目标类别多，锥桶、车辆、行人混合处理会降低锚点表达质量。

### 2. 技术方案概述

本发明提出一种面向港口异构交通参与体的转弯感知多模态运动预测方法，包括：

1. 从港口自动驾驶训练数据中读取目标类别、三维框航向角、未来轨迹和未来轨迹有效掩码；
2. 将未来轨迹变换到目标局部坐标系，使局部 +Y 方向表示目标当前朝向前方；
3. 基于路径长度、净位移、航向变化和横向偏移比例，将目标未来运动划分为静止/低速、直行、轻微转弯和大转弯；
4. 按目标类别构建锚点组，例如行人组、车辆/作业设备组、锥桶组；
5. 按运动桶分层保留固定数量的锚点槽位，必要时基于转弯强度对样本加权聚类，生成覆盖直行、低速、左右转弯和机动作业的多模态运动锚点；
6. 将锚点加载到多模态运动预测网络中，作为候选轨迹生成、初始参考轨迹或位置编码的先验信息；
7. 训练时计算每个候选预测模式相对真实轨迹的平均位移误差和终点位移误差，并采用二者的联合度量选择监督模式；
8. 训练时按运动桶对轨迹回归损失、终点误差项和/或模式分类损失进行加权，并对批次权重均值归一；
9. 推理时输出多个候选轨迹及置信度，以置信度最高轨迹作为最终输出，并通过最高置信度候选/最优候选诊断定位候选覆盖不足、模式评分错误或目标漏检问题。

### 3. 有益效果

1. 转弯感知锚点改善了候选轨迹覆盖，使车辆组在固定候选模式数量下包含真实左右转弯或作业式转向模板。
2. 平均位移误差与终点位移误差联合的模式分配使训练监督更关注终点误差，适合港口大车转弯和长距离轨迹预测。
3. 转弯分桶加权损失提高了轻微转弯和大转弯样本的训练影响力，并通过均值归一避免整体损失尺度漂移。
4. 锥桶独立分组降低近静止/抖动目标对车辆锚点的污染。
5. 最高置信度候选/最优候选诊断把问题拆成候选覆盖不足、模式评分错误、目标漏检三类，便于后续优化和专利实施例扩展。

## 三、具体实施方式

### 步骤 S1：采集和读取训练样本

训练样本包括港口连续帧传感器数据、目标三维框、目标类别、目标航向角、目标未来轨迹位置序列及其有效掩码。未来轨迹长度和候选模式数量可根据预测时域、模型容量和实时性要求配置；一个工程实施例中使用固定长度未来轨迹和固定数量候选模式。

### 步骤 S2：类别分组

{table_md(
    ["组号", "类别", "设计理由"],
    [
        ["G0", "行人", "行人尺度、速度和运动模式与车辆不同。"],
        ["G1", "小车、智能导引运输车、集卡、空/重载拖挂车、吊车、叉车及其他港口车辆/作业设备", "港口车辆或作业设备，存在直行、转弯、横移、低速作业等多种运动模式。"],
        ["G2", "锥桶", "锥桶大多近静止，单独分组可避免静态/抖动目标污染行人和车辆锚点。"],
    ],
)}

### 步骤 S3：目标局部坐标归一

对目标未来轨迹相对位移进行局部坐标变换，使局部 +Y 方向表示目标当前朝向前方。设目标当前航向角为 ψ，未来相对位移为 p，则局部轨迹点可表示为：

```text
p_local = R(π/2 - ψ) · p（1）
```

式（1）中，p_local 表示目标局部坐标系下的未来轨迹点或相对位移，ψ 表示目标当前航向角，p 表示场景坐标系下的未来相对位移，R(·) 表示二维旋转矩阵。

该变换使直行车辆未来轨迹主要沿局部 +Y 方向分布，左右转弯轨迹表现为向局部 X 方向逐渐偏移的弧线，从而有利于跨场景、跨朝向地聚类运动模式。

### 步骤 S4：计算运动形态特征与转弯强度

对局部轨迹计算路径长度 path、终点净位移 net、航向变化 Δθ 和横向偏移比例 lat_ratio：

```text
path = Σ ||p_t - p_(t-1)||
net = ||p_T||
Δθ = |wrapπ(atan2(v_last,y, v_last,x) - atan2(v_first,y, v_first,x))|
lat_ratio = max_t |p_t · n| / max(net, ε)（2）
```

式（2）中，path 表示路径长度，net 表示终点净位移，Δθ 表示航向变化，lat_ratio 表示横向偏移比例，p_t 表示第 t 个未来轨迹点，p_T 表示最后一个有效未来轨迹点，t 表示时间步索引，T 表示最后时间步，v_first 和 v_last 分别表示首个和最后一个有效运动段，n 表示终点方向的法向量，ε 表示用于避免分母为零的极小正数，wrapπ(·) 表示将角度归一化到 [-π, π] 或等价周期区间的函数，atan2(·) 表示用于计算方向角的反正切函数。若有效运动段长度过小，可将该样本视为静止/低速样本，以避免近静止抖动导致误判。

### 步骤 S5：运动分桶

一个实施例中，按如下规则分桶：

{table_md(
    ["运动桶", "示例判定条件", "作用"],
    [
        ["静止/低速", "路径长度或净位移低于预设阈值", "隔离静止/低速样本，避免稀释转弯效果。"],
        ["直行", "航向变化和横向偏移比例均低于直行阈值", "表示直行或近似直行运动。"],
        ["轻微转弯", "航向变化或横向偏移比例处于中等范围", "表示轻微转弯、横向偏移或低曲率变向。"],
        ["大转弯", "航向变化或横向偏移比例超过大转弯阈值", "表示大角度转弯、掉头、作业式转向等复杂运动。"],
    ],
)}

上述阈值为实施例，正式保护范围不应限于具体数值。

### 步骤 S6A：分层保留式锚点生成

一个优选实施例中，为车辆或作业设备类别预留若干运动族槽位，例如静止/低速、短直行、长直行、左向横移或转弯、右向横移或转弯、大机动转弯。

每个槽位从对应运动桶中选取均值中心、加权中心或代表性样本作为锚点；若某一槽位样本不足，可从相邻运动族回退补足，或使用该类别组的全局代表轨迹补足。该方式将固定数量的候选模式显式分配给港口作业中常见的运动族，能够避免普通聚类被高频直行/静止样本占据，从结构上保证左右转弯和急转弯在有限候选模式中获得表达。

### 步骤 S6B：转弯样本加权聚类生成锚点

为提升转弯样本在聚类中的影响力，根据转弯强度设置样本权重。一个实施例中：

```text
w(Δθ) = 1,                                   Δθ ≤ θ1
w(Δθ) = 1 + (α - 1)(Δθ - θ1)/(θ2 - θ1),     θ1 < Δθ < θ2
w(Δθ) = α,                                   Δθ ≥ θ2（3）
```

式（3）中，w(Δθ) 表示航向变化为 Δθ 时的样本权重，θ1、θ2 和 α 分别为低转弯阈值、高转弯阈值和最大权重系数，可根据数据集分布设置。随后在每个类别组内对展平后的未来轨迹执行带权聚类：

```text
C* = arg min_(C1,...,CM) Σ_i w_i · min_m ||τ_i - C_m||_2^2（4）
```

式（4）中，C* 表示最优锚点集合，τ_i 表示第 i 条未来轨迹样本，C_m 表示第 m 个锚点中心，w_i 表示第 i 条样本对应的转弯权重，M 表示锚点数量，i 表示样本索引，m 表示锚点索引。最大权重系数可以由训练集运动桶占比、转弯样本数量、验证集误差或业务安全权重确定。若权重过小，转弯长尾样本仍可能被直行样本淹没；若权重过大，可能放大异常轨迹、标注噪声或少量特殊作业样本。因此，在工程实施中可配合权重上限、样本最小数约束和批次均值归一，以兼顾转弯召回和整体稳定性。

### 步骤 S7：将锚点用于运动预测头

运动预测模块读取锚点数据，得到各类别组对应的多模态轨迹锚点。训练或推理时，先根据跟踪目标类别选择对应类别组的多个候选锚点，再根据目标当前三维框中心和航向角将局部锚点变换到场景坐标，作为多模态预测的初始参考轨迹和位置编码输入。

### 步骤 S7A：预测网络总体架构及转弯感知模块嵌入方式

本发明不限定运动预测网络必须采用特定骨干网络、特定层数或特定端到端框架。一个通用实施例中，预测网络包括感知特征输入模块、目标状态编码模块、场景交互编码模块、多模态轨迹解码模块、模式评分分支和轨迹回归分支。

感知特征输入模块用于接收点云、图像、多传感器融合特征或由上游检测/跟踪模块输出的目标历史状态；目标状态编码模块用于编码目标类别、当前位置、尺寸、航向角、历史运动状态等信息；场景交互编码模块用于建模目标之间以及目标与港区环境之间的交互关系；多模态轨迹解码模块基于上述特征和转弯感知锚点先验生成多个候选未来轨迹；模式评分分支输出各候选轨迹的置信度，轨迹回归分支输出各候选轨迹在未来时间步的位置偏移或绝对坐标。

转弯感知模块可嵌入在以下位置：其一，在训练前或训练过程中生成按类别组和运动族组织的锚点库；其二，在预测网络输入或解码阶段，根据目标类别、中心和航向角选择对应锚点并转换到场景坐标；其三，在训练损失阶段，根据真实未来轨迹所属运动桶选择监督模式并调整损失权重；其四，在评估阶段，比较最高置信度候选和最优候选，用于判断后续应优化候选覆盖、模式评分还是上游检测跟踪。

因此，本发明的保护重点不在于某一种特定网络骨干，而在于将目标局部坐标下的转弯度量、港口类别分组、分层转弯锚点、多模态轨迹解码和转弯分桶训练损失组合到运动预测网络中的方法。该组合可应用于基于卷积、注意力机制、图神经网络、Transformer 或多传感器融合的不同运动预测网络。

### 步骤 S8：平均位移误差与终点位移误差联合模式分配

对每个目标，模型输出多个候选未来轨迹及其置信度。训练时计算每个候选模式的平均位移误差和终点位移误差，并用下式选择监督模式：

```text
s_k = ADE_k + λ · FDE_k
k* = arg min_k s_k（5）
```

式（5）中，k 表示候选预测模式索引，s_k 表示第 k 个候选模式的联合误差，ADE_k 表示第 k 个候选模式的平均位移误差，FDE_k 表示第 k 个候选模式的终点位移误差，λ 表示终点误差权重，k* 表示被选为监督目标的候选模式。相比只用平均位移误差选择模式，该方式能让训练监督更关注终点落点，适合港口大车转弯、拖挂车长距离移动和作业式转向。

### 步骤 S9：转弯分桶样本加权损失

训练时根据真实未来轨迹所属运动桶设置样本权重。一个实施例中：

{table_md(
    ["运动桶", "回归/轨迹损失权重", "模式分类附加权重"],
    [
            ["静止/低速", "低于或等于基准权重", "基准权重"],
            ["直行", "基准权重", "基准权重"],
            ["轻微转弯", "高于基准权重", "高于基准权重"],
            ["大转弯", "高于轻微转弯权重", "最高分类权重"],
    ],
)}

损失可写为：

```text
L = β_cls · L_cls + β_reg · L_reg
  + β_fde · L_fde（6）
```

式（6）中，L 表示总损失，L_cls、L_reg 和 L_fde 分别表示模式分类损失、轨迹回归损失和终点误差辅助项，β_cls、β_reg 和 β_fde 分别表示对应损失项权重；其中 cls、reg 和 fde 分别对应分类项、轨迹回归项和终点误差项。不同运动桶的权重可按批次均值归一，以保持整体损失尺度稳定。具体权重数值为实施例，不限定保护范围。

### 步骤 S10：最高置信度候选/最优候选诊断与失效归因

评估时同时统计最优候选指标和最高置信度候选指标。最优候选指标指多个候选轨迹中事后按真实轨迹选择误差最小的一条，仅用于诊断；最高置信度候选指标指模型实际推理时置信度最高的候选轨迹。二者差异可用于定位：

{table_md(
    ["现象", "含义", "优化方向"],
    [
        ["最高置信度误差较大，最优候选误差较小", "候选轨迹里已有较优结果，但分类分数选错。", "优化模式评分、监督分配或分类损失加权。"],
        ["最高置信度误差和最优候选误差均较大", "候选本身覆盖不足。", "优化锚点生成或轨迹回归能力。"],
        ["目标未匹配", "运动预测模块未获得正确目标。", "优化目标检测或跟踪质量。"],
    ],
)}

## 四、创新点

1. 将港口运动预测问题拆成“候选覆盖”和“模式选择”两层，分别用转弯感知锚点和平均位移误差/终点位移误差联合分配及转弯加权训练解决。
2. 在目标局部坐标系下统一度量转弯强度，使不同朝向和不同港区位置的轨迹可比较。
3. 将港口车辆/作业设备、行人和锥桶分别成组，并在车辆/作业设备组内设置静止、短直行、长直行、左/右转弯和大机动转弯等运动族槽位。
4. 通过分层保留式锚点生成，在候选模式数量不增加的前提下为真实转弯候选预留表达空间；带权聚类可作为各槽位代表轨迹生成或补充实施方式。
5. 将锥桶等近静止目标独立成组，降低静态/抖动目标对车辆锚点的污染。
6. 训练时用平均位移误差与终点位移误差的联合度量选择监督模式，并按运动桶对轨迹损失和/或模式分类损失加权。
7. 通过批次均值归一控制转弯加权后的损失尺度，避免多任务训练中整体权重漂移。
8. 引入最高置信度候选/最优候选差异诊断体系，将转弯预测失败归因到候选覆盖、模式评分或目标漏检；该诊断体系可作为优化工具，不必将未充分验证的模式评分增强结果作为主要保护效果。

## 五、需要保护的要点

1. 一种从港口自动驾驶训练数据中提取未来轨迹，并根据目标类别、航向角和未来轨迹有效性生成运动锚点的方法。
2. 一种将未来轨迹变换到目标局部坐标系，并令局部 +Y 方向表示目标前方的轨迹归一方法。
3. 一种根据路径长度、净位移、航向变化和横向偏移比例划分静止/直行/轻微转弯/大转弯运动桶的方法。
4. 一种为固定数量的预测模式预留静止、短直行、长直行、左/右转弯和急转弯槽位的分层锚点生成方法。
5. 一种基于转弯强度设置样本权重，并使用带权聚类生成或补充分层锚点代表轨迹的方法。
6. 一种将行人、港口车辆类目标和锥桶分别分组生成运动锚点的方法。
7. 一种将转弯感知锚点加载到多模态运动预测网络，并根据目标类别、中心和航向变换到场景坐标后参与运动预测的方法。
8. 一种采用平均位移误差与终点位移误差联合度量选择多模态轨迹监督模式的训练方法。
9. 一种按转弯运动桶对轨迹回归损失、终点误差项和/或模式分类损失进行加权并归一化的方法。
10. 一种利用最高置信度候选与最优候选指标差异对运动预测失败进行候选覆盖不足、模式评分错误和目标漏检分类的方法。
11. 实现上述方法的装置、电子设备和计算机可读存储介质。

## 六、建议权利要求布局草案

### 独立权利要求1：方法

一种面向港口自动驾驶车辆的转弯感知多模态运动预测方法，包括：

1. 获取目标的类别、当前位姿、航向角、未来轨迹及未来轨迹有效掩码；
2. 将所述未来轨迹按照所述目标的航向角变换到目标局部坐标系；
3. 根据目标局部坐标系下的路径长度、净位移、航向变化和横向偏移比例确定目标未来运动所属的运动桶；
4. 根据目标类别和运动桶生成或选择多模态运动锚点；
5. 将多模态运动锚点输入运动预测网络，输出多个候选未来轨迹及其置信度；
6. 在训练阶段根据候选未来轨迹相对于真实未来轨迹的平均位移误差和终点位移误差确定监督模式，并根据所述运动桶对训练损失进行加权；
7. 在推理阶段根据候选未来轨迹置信度输出目标未来轨迹。

### 从属权利要求建议

1. 所述目标局部坐标系中，局部 +Y 方向表示目标当前朝向前方。
2. 所述运动桶包括静止/低速、直行、轻微转弯和大转弯。
3. 所述多模态运动锚点通过运动族槽位保留生成，所述运动族槽位至少包括静止/低速、短直行、长直行、左向横移或转弯、右向横移或转弯、大机动转弯中的多种。
4. 所述多模态运动锚点通过带转弯权重的聚类生成或补充，其中转弯权重根据航向变化、横向偏移比例或二者组合确定。
5. 所述目标类别至少包括行人组、车辆/作业设备组和锥桶组。
6. 所述航向变化和横向偏移比例分别与第一航向阈值、第二航向阈值、第一横向偏移阈值和第二横向偏移阈值比较，以确定直行、轻微转弯和大转弯。
7. 所述监督模式根据平均位移误差与终点位移误差联合度量的最小值确定。
8. 所述训练损失包括模式分类损失、轨迹回归损失和终点误差辅助项中的至少一种。
9. 所述训练损失权重按运动桶分别设置，且转弯运动桶权重高于直行运动桶权重。
10. 所述训练损失权重按批次均值归一，以抑制不同训练批次中转弯样本占比变化导致的损失尺度漂移。
11. 所述方法还包括基于最高置信度候选误差与最优候选误差的差异输出失效归因结果。
12. 所述方法还包括根据候选轨迹置信度与真实轨迹误差之间的差异，对模式分类损失进行增强。

### 独立权利要求2：装置

一种面向港口自动驾驶车辆的转弯感知多模态运动预测装置，包括：数据获取模块、局部坐标变换模块、运动桶划分模块、锚点生成或选择模块、运动预测模块、训练损失构建模块和推理输出模块；其中各模块用于执行独立权利要求1所述方法中的对应步骤。

### 独立权利要求3：电子设备

一种电子设备，包括处理器和存储器，所述存储器中存储有计算机程序，所述计算机程序在由所述处理器执行时实现独立权利要求1所述的方法。

### 独立权利要求4：计算机可读存储介质

一种计算机可读存储介质，其上存储有计算机程序，所述计算机程序被处理器执行时实现独立权利要求1所述的方法。

## 七、实验数据与效果

### 1. 阶段性整体消融结果

表7-1给出各实验组设置。表7-2给出阶段性整体消融结果。指标中，minADE、minFDE 和 MR 越低越好，AMOTA 和 NDS 越高越好；minADE 和 minFDE 的单位为米。

表7-1：实验组设置

{experiment_setup_table}

表7-2：阶段性整体消融结果

{overall_summary_table}

结论：本表为阶段性消融结果，用于说明方案方向和各模块贡献，不能替代完整训练收敛后的最终性能结论。当前结果显示：单独更换转弯感知锚点后终点误差略有下降；转弯分桶训练目标是当前运动预测单项指标最优配置；转弯感知锚点与转弯分桶训练目标组合后，运动预测指标略弱于单独转弯分桶训练目标，但综合感知跟踪指标更好。正式提交或对外披露前，建议补充更长训练轮次或验证集最佳检查点结果，以确认趋势稳定性。

### 2. 转弯分桶结果

表7-3聚焦轻微转弯和大转弯样本，并给出相对于实验组A的变化。该表用于突出与本发明技术问题直接相关的转弯场景效果。

表7-3：转弯场景重点指标与相对变化

{turn_focus_table}

结论：转弯加权训练对轻微转弯和大转弯的终点误差改善更明显。例如实验组C相对于实验组A，轻微转弯minFDE和大转弯minFDE均有明显下降。仅靠锚点的收益有限，尤其大转弯仍会暴露模式选择和分类评分问题，因此本交底书将训练目标和模式分配纳入核心方案。实验组E用于验证“候选存在但最高置信度未选中”的模式评分增强方向；现有阶段性结果只显示大转弯MR略有改善，尚未形成稳定、全面的性能提升，因此不宜把该实验作为主要效果宣称，应将其作为诊断体系导出的后续优化实施例。

### 3. 视觉复查结论

代表样本显示：部分大转弯场景中，本方案能够给出更接近真实轨迹的候选；但也存在最优候选误差较小而最高置信度轨迹误差较大的样本，说明当前瓶颈已经从“有没有候选”部分转向“能不能把正确候选排到最高置信度”。

## 八、附图及附图说明

1. 图1：转弯感知运动锚点生成流程。该图展示训练数据读取、港口类别分组、目标局部坐标归一、运动形态特征计算、运动族槽位保留、样本加权和锚点输出之间的关系。其重点在于说明锚点不是从全量轨迹直接普通聚类得到，而是先结合港口类别和转弯语义进行结构化约束。
2. 图2：普通运动锚点与转弯感知运动锚点的形状对比。该图用于展示普通锚点容易集中在直行或低曲率轨迹，而转弯感知锚点在有限候选模式下保留左右转弯和大机动转弯形状。图中差异体现了分层锚点和转弯加权对候选覆盖的影响。
3. 图3A至图3C：转弯目标聚焦可视化对比示例。三图分别选取卡车轻微转弯、起重机大角度转弯等不同目标样本；每张图左侧为普通锚点方案的最高置信度轨迹，右侧为本方案的最高置信度轨迹，白色曲线为真实未来轨迹，坐标轴分别表示目标局部视角下的横向位置和前向位置，单位为米。图3A中普通方案最高置信度终点误差为8.06米，本方案为0.26米；图3B中普通方案为6.53米，本方案为0.81米；图3C中普通方案为7.71米，本方案为1.18米。上述图仅作为定性示例，与前述分桶统计结果共同说明本方案对转弯目标的轨迹贴合程度。
4. 图4：转弯感知损失与模式分配流程。该图展示平均位移误差/终点位移误差联合选择监督模式，以及基于运动桶的样本加权和批次均值归一。其重点在于说明本发明同时处理“监督哪个候选模式”和“转弯样本训练权重不足”两个问题。
5. 图5：最高置信度候选/最优候选模式诊断示意。该图展示候选存在但置信度选错、候选覆盖不足和目标漏检三类失效模式。该诊断可用于决定后续应优化锚点生成、轨迹回归、模式评分还是上游检测跟踪。
6. 图6：预测网络总体架构及转弯感知模块嵌入位置示意。该图展示传感器特征和目标历史状态经目标/场景编码、交互编码后进入多模态轨迹解码器；转弯感知锚点库根据目标类别、中心和航向角选择并转换为场景坐标，作为解码器的锚点先验；训练阶段结合真实未来轨迹和运动桶标签进行联合模式分配及运动桶加权，推理阶段输出候选轨迹及其置信度。该图用于说明本方案可嵌入不同运动预测骨干网络，而不限定具体网络层数或代码实现。

## 九、可扩展实施例

1. 转弯权重函数可采用分段线性、指数函数、sigmoid 函数或查表方式。
2. 转弯强度可结合曲率、横向偏移比、终点方向变化、历史角速度等多种运动特征。
3. 类别分组可根据港口设备类型调整，例如将吊车、叉车、拖挂车分别设置独立锚点组。
4. 该方法可用于点云、视觉、多传感器融合等不同感知预测模型。
5. 模式评分可进一步采用温度缩放、排序损失、最高置信度候选/最优候选差异损失或困难样本挖掘。
6. 分层锚点槽位数量可根据预测时域、实时性约束和目标类别调整，不限于六个槽位。
7. 运动桶权重可与课程学习结合，在训练早期采用较低转弯权重，在训练后期逐步提高转弯长尾样本权重。
8. 锚点可基于训练过程中的困难样本或周期性验证集误差动态更新，以适应港区路线、作业工况或设备类型变化。
9. 模式分类可引入对比学习、排序学习或候选间间隔约束，使接近真实轨迹的候选获得更高置信度。
10. 失效归因结果可反馈到数据采样、场景挖掘和标注复核流程，用于补充大转弯、倒车、横移和装卸作业等长尾场景。
11. 转弯分桶可结合道路拓扑、作业区域属性或任务指令作为附加条件，但这些附加条件不是本方案实施的必要条件。
12. 本方法可与不同轨迹预测骨干网络结合，包括基于注意力机制、图结构交互或多传感器融合的预测网络。

## 十、附图生成说明

本交底书附图均为技术流程示意或内部实验可视化摘录。正式提交时可保留与技术方案直接相关的流程图和对比图；完整的生成命令、代码路径和实验目录已另存为内部备查文件，不放入交底书正文。
"""
    return text


def build_repro_markdown() -> str:
    return f"""# 转弯感知运动预测交底书内部复现记录

该文件仅用于内部复现、数据追溯和代理人答疑，不建议作为交底书正文提交。

## 1. 生成命令

```bash
python3 documents/patent_2026_motion/build_turnaware_patent_docx_v05.py
```

## 2. 文本来源

{table_md(
    ["来源", "用途"],
    [
        ["研发总结文档", "最新实验结论、消融结果、问题归因。"],
        ["上一版交底书草案", "上一版交底书结构、背景、基础实施方式。"],
        ["轨迹损失实现文件", "平均位移误差/终点位移误差联合模式分配、转弯分桶权重、分类权重与运动桶分类实现。"],
        ["锚点生成脚本", "带权聚类的转弯感知锚点生成。"],
        ["分层锚点生成脚本", "分层保留式锚点生成备选实施例。"],
        ["转弯分桶评估脚本", "转弯分桶评估和最高置信度候选/最优候选诊断。"],
        ["实验配置文件", "实验配置和损失参数。"],
    ],
)}

## 3. 实验数据源

{table_md(
    ["文件", "用途"],
    [
        ["sources/v0.5_epoch6_overall_ablation_summary.csv", "交底书整体消融结果表。"],
        ["sources/v0.5_turn_bucket_epoch6_summary.csv", "交底书转弯分桶结果表。"],
        ["原始转弯分桶评估结果", "原始转弯分桶指标。"],
        ["代表性可视化案例清单", "代表可视化案例和最高置信度候选/最优候选诊断样本。"],
    ],
)}

## 4. 图片生成/复用说明

{table_md(
    ["图片", "生成或复用方法", "数据源"],
    [
        ["fig1_turnaware_anchor_generation_flow.png", "由本文档生成脚本使用 PIL 绘制。", "算法流程来源：类别分组、局部坐标归一、运动分桶、分层锚点与转弯加权方案。"],
        ["fig2_anchor_compare.png", "由本文档生成脚本读取锚点 pkl 并使用 PIL 直接绘制中文图。", "数据源：sources/motion_anchor_infos_kl.pkl 与 sources/motion_anchor_infos_kl_turnaware_3grp.pkl；绘图保留原始锚点坐标数据。"],
        ["fig3a_turn_focus_truck_1522.png、fig3b_turn_focus_crane_204.png、fig3c_turn_focus_crane_2409.png", "由 tools/analysis_tools/visualize_lidar_motion_focus_compare.py 使用 --chinese-labels、--hide-frame-title、--show-axis-labels、--patent-layout 与 --large-text 重新渲染样本索引 1522、204、2409 后复制。", "数据源：base_e2e_lidar/eval_epoch6_results.pkl、base_e2e_lidar_turnaware_turnloss/eval_epoch6_results.pkl、data/kl_8/kl_infos_val.pkl 及对应点云；输出目录：sources/fig3_focus_patent_ttl_zh。"],
        ["fig4_turnaware_loss_mode_assignment_flow.png", "由本文档生成脚本使用 PIL 绘制。", "算法流程和配置参数：轨迹损失实现文件、实验配置文件。"],
        ["fig5_candidate_confidence_diagnosis.png", "由本文档生成脚本使用 PIL 绘制。", "诊断逻辑来源：转弯分桶评估脚本和研发总结文档。"],
        ["fig6_prediction_network_architecture.png", "由本文档生成脚本使用 PIL 绘制。", "内容来源：本交底书中的通用预测网络架构描述，用于说明转弯感知锚点、轨迹解码、训练损失和推理输出的嵌入位置。"],
    ],
)}

图3重生成命令：

```bash
PYTHONPATH="$PWD:$PYTHONPATH" /home/baojiali/anaconda3/envs/uniad_train/bin/python tools/analysis_tools/visualize_lidar_motion_focus_compare.py \
  --config projects/configs/stage2_e2e_lidar/base_e2e_lidar.py \
  --result base projects/work_dirs/stage2_e2e_lidar/base_e2e_lidar/eval_epoch6_results.pkl \
  --result turnaware_turnloss projects/work_dirs/stage2_e2e_lidar/base_e2e_lidar_turnaware_turnloss/eval_epoch6_results.pkl \
  --case 1522:Truck:10:base_top1_8.06_vs_scheme_0.26 \
  --case 204:Crane:5:base_top1_6.53_vs_scheme_0.81 \
  --case 2409:Crane:4:base_top1_7.71_vs_scheme_1.18 \
  --out-dir documents/patent_2026_motion/sources/fig3_focus_patent_ttl_zh \
  --split val \
  --score-thr 0.25 \
  --point-stride 5 \
  --zoom-margin 12 \
  --chinese-labels \
  --hide-frame-title \
  --show-axis-labels \
  --patent-layout \
  --large-text
```

## 5. 完整复现 JSON

```text
documents/patent_2026_motion/sources/v0.5_reproducibility_record.json
```
"""


def turn_models_for_doc():
    return [
        ("base_e2e_lidar", "old anchor + original loss"),
        ("base_e2e_lidar_turnaware", "turn-aware anchor + original loss"),
        ("base_e2e_lidar_turnloss", "old anchor + turn-aware loss"),
        ("base_e2e_lidar_turnaware_turnloss", "turn-aware anchor + turn-aware loss"),
        ("base_e2e_lidar_turnaware_modescore", "turn-aware anchor + turn-aware loss + turn cls score"),
    ]


def set_run_font(run, size=11, bold=False, italic=False, font_name="宋体"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def set_cell_text(cell, text, bold=False, size=10.5, font_name="宋体", align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold, font_name=font_name)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="9CA3AF", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(int(width_cm * 567)))
    width.set(qn("w:type"), "dxa")


def format_table(table, header=True, first_col=False, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            set_cell_border(cell)
            if widths and col_idx < len(widths):
                set_cell_width(cell, widths[col_idx])
            if header and row_idx == 0:
                shade_cell(cell, "EAF0F7")
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_run_font(run, size=10.5, bold=True, font_name="黑体")
            elif first_col and col_idx == 0:
                shade_cell(cell, "F5F7FB")


def add_doc_paragraph(doc, text="", size=11, bold=False, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def is_formula_block(code_lines):
    text = "\n".join(line.strip() for line in code_lines)
    if not text:
        return False
    formula_markers = [
        "=", "Σ", "||", "Δ", "θ", "λ", "ε", "arg min", "max_", "min_",
        "·", "≤", "≥", "C*", "w(", "p_local", "lat_ratio", "总损失", "联合误差",
    ]
    non_formula_markers = [
        "python3 ", ".json", ".md", ".docx", ".csv", ".png",
        "静止/低速、短直行", "短直行、长直行",
    ]
    return any(marker in text for marker in formula_markers) and not any(marker in text for marker in non_formula_markers)


FORMULA_RENDERINGS = {
    "p_local = R(π/2 - ψ) · p（1）": [
        [("p", None, None), ("local", "sub", None), (" = R(π/2 − ψ) · p    （1）", None, None)],
    ],
    "path = Σ ||p_t - p_(t-1)||\nnet = ||p_T||\nΔθ = |wrapπ(atan2(v_last,y, v_last,x) - atan2(v_first,y, v_first,x))|\nlat_ratio = max_t |p_t · n| / max(net, ε)（2）": [
        [("path = Σ", None, None), ("t", "sub", None), (" ||p", None, None), ("t", "sub", None), (" − p", None, None), ("t−1", "sub", None), ("||", None, None)],
        [("net = ||p", None, None), ("T", "sub", None), ("||", None, None)],
        [("Δθ = |wrap", None, None), ("π", "sub", None), ("(θ", None, None), ("last", "sub", None), (" − θ", None, None), ("first", "sub", None), (")|", None, None)],
        [("lat", None, None), ("ratio", "sub", None), (" = max", None, None), ("t", "sub", None), (" |p", None, None), ("t", "sub", None), (" · n| / max(net, ε)    （2）", None, None)],
    ],
    "w(Δθ) = 1,                                   Δθ ≤ θ1\nw(Δθ) = 1 + (α - 1)(Δθ - θ1)/(θ2 - θ1),     θ1 < Δθ < θ2\nw(Δθ) = α,                                   Δθ ≥ θ2（3）": [
        [("w(Δθ) = 1,    Δθ ≤ θ", None, None), ("1", "sub", None)],
        [("w(Δθ) = 1 + (α − 1)(Δθ − θ", None, None), ("1", "sub", None), (")/(θ", None, None), ("2", "sub", None), (" − θ", None, None), ("1", "sub", None), ("),    θ", None, None), ("1", "sub", None), (" < Δθ < θ", None, None), ("2", "sub", None)],
        [("w(Δθ) = α,    Δθ ≥ θ", None, None), ("2", "sub", None), ("    （3）", None, None)],
    ],
    "C* = arg min_(C1,...,CM) Σ_i w_i · min_m ||τ_i - C_m||_2^2（4）": [
        [("C", None, None), ("*", "super", None), (" = arg min", None, None), ("C₁,...,Cₘ", "sub", None), (" Σ", None, None), ("i", "sub", None), (" w", None, None), ("i", "sub", None), (" · min", None, None), ("m", "sub", None), (" ||τ", None, None), ("i", "sub", None), (" − C", None, None), ("m", "sub", None), ("||", None, None), ("2", "super", None), ("    （4）", None, None)],
    ],
    "s_k = ADE_k + λ · FDE_k\nk* = arg min_k s_k（5）": [
        [("sₖ = ADEₖ + λ · FDEₖ", None, None)],
        [("k* = arg minₖ sₖ    （5）", None, None)],
    ],
    "L = β_cls · L_cls + β_reg · L_reg\n+ β_fde · L_fde（6）": [
        [("L = β_cls · L_cls + β_reg · L_reg", None, None)],
        [("  + β_fde · L_fde    （6）", None, None)],
    ],
}

FORMULA_OMML_TEXTS = {
    "p_local = R(π/2 - ψ) · p（1）": [
        "p_local = R(π/2 − ψ) · p    （1）",
    ],
    "path = Σ ||p_t - p_(t-1)||\nnet = ||p_T||\nΔθ = |wrapπ(atan2(v_last,y, v_last,x) - atan2(v_first,y, v_first,x))|\nlat_ratio = max_t |p_t · n| / max(net, ε)（2）": [
        "path = Σ_t ||p_t − p_(t−1)||",
        "net = ||p_T||",
        "Δθ = |wrapπ(θ_last − θ_first)|",
        "lat_ratio = max_t |p_t · n| / max(net, ε)    （2）",
    ],
    "w(Δθ) = 1,                                   Δθ ≤ θ1\nw(Δθ) = 1 + (α - 1)(Δθ - θ1)/(θ2 - θ1),     θ1 < Δθ < θ2\nw(Δθ) = α,                                   Δθ ≥ θ2（3）": [
        "w(Δθ) = 1,    Δθ ≤ θ_1",
        "w(Δθ) = 1 + (α − 1)(Δθ − θ_1)/(θ_2 − θ_1),    θ_1 < Δθ < θ_2",
        "w(Δθ) = α,    Δθ ≥ θ_2    （3）",
    ],
    "C* = arg min_(C1,...,CM) Σ_i w_i · min_m ||τ_i - C_m||_2^2（4）": [
        "C* = arg min_(C₁,...,C_M) Σ_i w_i · min_m ||τ_i − C_m||₂²    （4）",
    ],
    "s_k = ADE_k + λ · FDE_k\nk* = arg min_k s_k（5）": [
        "s_k = ADE_k + λ · FDE_k",
        "k* = arg min_k s_k    （5）",
    ],
    "L = β_cls · L_cls + β_reg · L_reg\n+ β_fde · L_fde（6）": [
        "L = β_cls · L_cls + β_reg · L_reg",
        "+ β_fde · L_fde    （6）",
    ],
}


def add_formula_run(paragraph, text, script=None):
    run = paragraph.add_run(text)
    set_run_font(run, size=12, font_name="Cambria Math")
    if script:
        vert_align = OxmlElement("w:vertAlign")
        vert_align.set(qn("w:val"), script)
        run._element.get_or_add_rPr().append(vert_align)
        run.font.size = Pt(8)
    return run


SCRIPT_TOKEN_RE = re.compile(r"([A-Za-z\u0370-\u03FFΣ]+)_\(([^)]*)\)|([A-Za-z\u0370-\u03FFΣ]+)_([A-Za-z0-9\u0370-\u03FF]+)")


def omml_run(text: str):
    mr = OxmlElement("m:r")
    mt = OxmlElement("m:t")
    mt.text = text
    mr.append(mt)
    return mr


def omml_subscript(base: str, sub: str):
    elem = OxmlElement("m:sSub")
    base_elem = OxmlElement("m:e")
    sub_elem = OxmlElement("m:sub")
    base_elem.append(omml_run(base))
    sub_elem.append(omml_run(sub))
    elem.append(base_elem)
    elem.append(sub_elem)
    return elem


def append_omml_expression(omath, text: str) -> None:
    pos = 0
    for match in SCRIPT_TOKEN_RE.finditer(text):
        if match.start() > pos:
            omath.append(omml_run(text[pos:match.start()]))
        if match.group(1) is not None:
            base, sub = match.group(1), match.group(2)
        else:
            base, sub = match.group(3), match.group(4)
        omath.append(omml_subscript(base, sub))
        pos = match.end()
    if pos < len(text):
        omath.append(omml_run(text[pos:]))


def split_equation_number(text: str):
    match = re.search(r"\s*[（(](\d+)[）)]\s*$", text)
    if not match:
        return text.rstrip(), ""
    body = text[: match.start()].rstrip()
    return body, f"({match.group(1)})"


def add_omath_paragraph(doc, text: str):
    text, number = split_equation_number(text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    omath = OxmlElement("m:oMath")
    append_omml_expression(omath, text)
    p._p.append(omath)
    if number:
        run = p.add_run(f"   {number}")
        set_run_font(run, size=11, font_name="宋体")
    return p


def add_formula_parts(paragraph, line_parts):
    for text, script, _ in line_parts:
        add_formula_run(paragraph, text, script=script)


def add_doc_formula_block(doc, code_lines):
    key = "\n".join(line.strip() for line in code_lines)
    texts = FORMULA_OMML_TEXTS.get(key)
    if texts is None:
        rendered = FORMULA_RENDERINGS.get(key)
        if rendered is None:
            texts = [line.strip().replace(" - ", " − ") for line in code_lines]
        else:
            texts = ["".join(text for text, _, _ in line_parts) for line_parts in rendered]
    for text in texts:
        add_omath_paragraph(doc, text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_doc_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True, font_name="黑体")


def add_doc_table(doc, headers, rows, widths=None, first_col=False):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, font_name="黑体", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for i, item in enumerate(row):
            set_cell_text(cells[i], item)
    format_table(table, header=True, first_col=first_col, widths=widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_md_table(lines, idx):
    headers = [x.strip() for x in lines[idx].strip().strip("|").split("|")]
    rows = []
    idx += 2
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        rows.append([x.strip() for x in lines[idx].strip().strip("|").split("|")])
        idx += 1
    return headers, rows, idx


def build_docx(md_text: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("坤浪科技专利技术交底书")
    set_run_font(r, size=20, bold=True, font_name="黑体")

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("# "):
            i += 1
            continue
        if stripped.startswith("## "):
            add_doc_heading(doc, stripped[3:], level=1)
        elif stripped.startswith("### "):
            add_doc_heading(doc, stripped[4:], level=2)
        elif stripped.startswith("|"):
            headers, rows, i = parse_md_table(lines, i)
            add_doc_table(doc, headers, rows)
            continue
        elif stripped.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            if is_formula_block(code):
                add_doc_formula_block(doc, code)
            else:
                table = doc.add_table(rows=1, cols=1)
                cell = table.rows[0].cells[0]
                set_cell_text(cell, "\n".join(code), size=9.5, font_name="Consolas")
                shade_cell(cell, "F3F4F6")
                set_cell_border(cell, color="D1D5DB")
        elif stripped.startswith("> "):
            table = doc.add_table(rows=1, cols=1)
            cell = table.rows[0].cells[0]
            set_cell_text(cell, stripped[2:], size=10.5)
            shade_cell(cell, "FFF7ED")
            set_cell_border(cell, color="F59E0B")
        elif stripped.startswith("- ") or stripped[:2].isdigit() and stripped[2:4] == ". " or stripped[:1].isdigit() and stripped[1:3] == ". ":
            add_doc_paragraph(doc, stripped, first_line=False)
        else:
            add_doc_paragraph(doc, stripped)
        i += 1

    add_doc_heading(doc, "附图", level=1)
    figures = [
        ("图1：转弯感知运动锚点生成流程", "fig1_turnaware_anchor_generation_flow.png"),
        ("图2：普通运动锚点与转弯感知运动锚点形状对比", "fig2_anchor_compare.png"),
        ("图3A：卡车轻微转弯目标聚焦对比", "fig3a_turn_focus_truck_1522.png"),
        ("图3B：起重机大角度转弯目标聚焦对比", "fig3b_turn_focus_crane_204.png"),
        ("图3C：起重机大角度转弯目标聚焦对比", "fig3c_turn_focus_crane_2409.png"),
        ("图4：转弯感知损失与模式分配流程", "fig4_turnaware_loss_mode_assignment_flow.png"),
        ("图5：最高置信度候选/最优候选诊断示意", "fig5_candidate_confidence_diagnosis.png"),
        ("图6：预测网络总体架构及转弯感知模块嵌入位置示意", "fig6_prediction_network_architecture.png"),
    ]
    figure_widths = {
        "fig2_anchor_compare.png": Inches(6.55),
        "fig3a_turn_focus_truck_1522.png": Inches(6.55),
        "fig3b_turn_focus_crane_204.png": Inches(6.55),
        "fig3c_turn_focus_crane_2409.png": Inches(6.55),
        "fig6_prediction_network_architecture.png": Inches(6.55),
    }
    for caption, filename in figures:
        path = FIG_DIR / filename
        if not path.exists():
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=figure_widths.get(filename, Inches(6.2)))
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        set_run_font(r, size=10)

    doc.save(OUT_DOCX)


def main() -> None:
    generate_figures()
    md_text = build_markdown()
    repro_text = build_repro_markdown()
    OUT_MD.write_text(md_text, encoding="utf-8")
    OUT_REPRO_MD.write_text(repro_text, encoding="utf-8")
    build_docx(md_text)
    if REPRO_JSON.exists():
        data = json.loads(REPRO_JSON.read_text(encoding="utf-8"))
        data["actual_outputs"] = [
            str(OUT_MD.relative_to(ROOT_DIR)),
            str(OUT_DOCX.relative_to(ROOT_DIR)),
            str(OUT_REPRO_MD.relative_to(ROOT_DIR)),
        ]
        REPRO_JSON.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUT_MD)
    print(OUT_DOCX)
    print(OUT_REPRO_MD)


if __name__ == "__main__":
    main()
